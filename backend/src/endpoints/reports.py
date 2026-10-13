"""Report generation endpoints for SIPGURU - daily/weekly/monthly with PDF/Excel export."""

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from datetime import datetime, timedelta
from typing import Optional
import io
import json

from src.database import get_db
from src.deps import get_current_user, require_admin

router = APIRouter()
db = get_db()


def _get_date_range(period: str, year: int, month: int, week: int = None):
    """Calculate date range based on period type."""
    if period == "daily":
        try:
            date = datetime(year, month, 1)
        except ValueError:
            date = datetime.now()
        start = date.replace(day=1).strftime("%Y-%m-%d")
        end = date.strftime("%Y-%m-%d")
        return start, end
    elif period == "weekly":
        # Week number (1-4) or current week
        if week is None:
            week = (datetime.now().day - 1) // 7 + 1
        start_date = datetime(year, month, 1) + timedelta(weeks=week - 1)
        end_date = start_date + timedelta(days=6)
        # Don't exceed month end
        if end_date.month != month:
            end_date = datetime(year, month + 1, 1) - timedelta(days=1)
        return start_date.strftime("%Y-%m-%d"), end_date.strftime("%Y-%m-%d")
    elif period == "monthly":
        try:
            start = datetime(year, month, 1).strftime("%Y-%m-%d")
            if month == 12:
                end = datetime(year + 1, 1, 1) - timedelta(days=1)
            else:
                end = datetime(year, month + 1, 1) - timedelta(days=1)
            return start, end.strftime("%Y-%m-%d")
        except ValueError:
            now = datetime.now()
            return now.replace(day=1).strftime("%Y-%m-%d"), now.strftime("%Y-%m-%d")
    else:
        now = datetime.now()
        return now.replace(day=1).strftime("%Y-%m-%d"), now.strftime("%Y-%m-%d")


@router.get("/attendance")
async def get_attendance_report(
    period: str = Query("daily", enum=["daily", "weekly", "monthly"]),
    year: int = Query(None),
    month: int = Query(None),
    week: int = Query(None),
    teacher_id: Optional[str] = Query(None),
    auth_user: dict = Depends(get_current_user),
):
    """Get attendance report for specified period with optional teacher filter."""
    if year is None:
        year = datetime.now().year
    if month is None:
        month = datetime.now().month
    
    start_date, end_date = _get_date_range(period, year, month, week)
    
    conn = db._get_connection()
    cursor = conn.cursor()
    
    # Build query based on filters
    query = """
        SELECT 
            al.id,
            al.teacher_id,
            u.name as teacher_name,
            u.nip,
            al.room_id,
            r.name as room_name,
            r.room_code,
            al.status,
            al.lateness_minutes,
            al.geolocation,
            al.created_at
        FROM attendance_logs al
        JOIN users u ON al.teacher_id = u.id
        LEFT JOIN rooms r ON al.room_id = r.id
        WHERE strftime('%Y-%m-%d', al.created_at) BETWEEN ? AND ?
    """
    params = [start_date, end_date]
    
    if teacher_id:
        query += " AND al.teacher_id = ?"
        params.append(teacher_id)
    
    query += " ORDER BY al.created_at DESC"
    
    cursor.execute(query, params)
    rows = cursor.fetchall()
    conn.close()
    
    # Parse geolocation JSON
    reports = []
    for row in rows:
        record = dict(row)
        if record.get("geolocation"):
            try:
                record["geolocation"] = json.loads(record["geolocation"])
            except:
                record["geolocation"] = None
        reports.append(record)
    
    # Calculate summary stats
    total = len(reports)
    on_time = sum(1 for r in reports if r["status"] == "on_time")
    late = sum(1 for r in reports if r["status"] == "late")
    avg_lateness = sum(r.get("lateness_minutes") or 0 for r in reports if r["status"] == "late") / max(late, 1)
    
    return {
        "period": period,
        "year": year,
        "month": month,
        "week": week,
        "start_date": start_date,
        "end_date": end_date,
        "summary": {
            "total": total,
            "on_time": on_time,
            "late": late,
            "attendance_rate": round(on_time / max(total, 1) * 100, 2),
            "avg_lateness_minutes": round(avg_lateness, 1)
        },
        "records": reports
    }


@router.get("/attendance/csv")
async def export_attendance_csv(
    period: str = Query("daily", enum=["daily", "weekly", "monthly"]),
    year: int = Query(None),
    month: int = Query(None),
    week: int = Query(None),
    teacher_id: Optional[str] = Query(None),
    auth_user: dict = Depends(get_current_user),
):
    """Export attendance report as CSV."""
    # Get data
    report = await get_attendance_report(
        period=period, year=year, month=month, week=week, 
        teacher_id=teacher_id, auth_user=auth_user
    )
    
    # Generate CSV
    output = io.StringIO()
    
    # Header
    output.write("Tanggal,Waktu,NIP,Nama Guru,Ruangan,Kode Ruangan,Status,Keterlambatan (menit)\n")
    
    # Data rows
    for record in report["records"]:
        created = record.get("created_at", "")
        date = created[:10] if created else ""
        time = created[11:16] if len(created) > 11 else ""
        output.write(
            f"{date},{time},{record.get('nip','')},{record.get('teacher_name','')},"
            f"{record.get('room_name','')},{record.get('room_code','')},"
            f"{record.get('status','')},{record.get('lateness_minutes') or ''}\n"
        )
    
    # Summary
    output.write(f"\nRingkasan:\n")
    output.write(f"Periode,{report['period']}\n")
    output.write(f"Total,{report['summary']['total']}\n")
    output.write(f"Tepat Waktu,{report['summary']['on_time']}\n")
    output.write(f"Terlambat,{report['summary']['late']}\n")
    output.write(f"Rate Kehadiran,{report['summary']['attendance_rate']}%\n")
    output.write(f"Rata-rata Keterlambatan,{report['summary']['avg_lateness_minutes']} menit\n")
    
    csv_content = output.getvalue()
    output.close()
    
    filename = f"laporan_absensi_{report['start_date']}_to_{report['end_date']}.csv"
    
    return StreamingResponse(
        io.BytesIO(csv_content.encode('utf-8')),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/attendance/excel")
async def export_attendance_excel(
    period: str = Query("daily", enum=["daily", "weekly", "monthly"]),
    year: int = Query(None),
    month: int = Query(None),
    week: int = Query(None),
    teacher_id: Optional[str] = Query(None),
    auth_user: dict = Depends(get_current_user),
):
    """Export attendance report as Excel (.xlsx)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed. Run: pip install openpyxl")
    
    # Get data
    report = await get_attendance_report(
        period=period, year=year, month=month, week=week,
        teacher_id=teacher_id, auth_user=auth_user
    )
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Absensi"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="064e3b", end_color="064e3b", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    # Headers
    headers = ["Tanggal", "Waktu", "NIP", "Nama Guru", "Ruangan", "Kode", "Status", "Keterlambatan (menit)"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Data rows
    for row_idx, record in enumerate(report["records"], 2):
        created = record.get("created_at", "")
        date = created[:10] if created else ""
        time = created[11:16] if len(created) > 11 else ""
        
        ws.cell(row=row_idx, column=1, value=date).border = thin_border
        ws.cell(row=row_idx, column=2, value=time).border = thin_border
        ws.cell(row=row_idx, column=3, value=record.get("nip", "")).border = thin_border
        ws.cell(row=row_idx, column=4, value=record.get("teacher_name", "")).border = thin_border
        ws.cell(row=row_idx, column=5, value=record.get("room_name", "")).border = thin_border
        ws.cell(row=row_idx, column=6, value=record.get("room_code", "")).border = thin_border
        ws.cell(row=row_idx, column=7, value=record.get("status", "")).border = thin_border
        ws.cell(row=row_idx, column=8, value=record.get("lateness_minutes")).border = thin_border
    
    # Summary section
    summary_row = len(report["records"]) + 3
    ws.cell(row=summary_row, column=1, value="RINGKASAN").font = Font(bold=True)
    
    summary_data = [
        ("Periode:", report["period"]),
        ("Tanggal:", f"{report['start_date']} - {report['end_date']}"),
        ("Total Absensi:", report["summary"]["total"]),
        ("Tepat Waktu:", report["summary"]["on_time"]),
        ("Terlambat:", report["summary"]["late"]),
        ("Rate Kehadiran:", f"{report['summary']['attendance_rate']}%"),
        ("Rata-rata Keterlambatan:", f"{report['summary']['avg_lateness_minutes']} menit"),
    ]
    
    for i, (label, value) in enumerate(summary_data, 0):
        ws.cell(row=summary_row + i + 1, column=1, value=label).font = Font(bold=True)
        ws.cell(row=summary_row + i + 1, column=2, value=value)
    
    # Column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 18
    
    # Save to buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"laporan_absensi_{report['start_date']}_to_{report['end_date']}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/attendance/pdf")
async def export_attendance_pdf(
    period: str = Query("daily", enum=["daily", "weekly", "monthly"]),
    year: int = Query(None),
    month: int = Query(None),
    week: int = Query(None),
    teacher_id: Optional[str] = Query(None),
    auth_user: dict = Depends(get_current_user),
):
    """Export attendance report as PDF."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
    except ImportError:
        raise HTTPException(status_code=500, detail="reportlab not installed. Run: pip install reportlab")
    
    # Get data
    report = await get_attendance_report(
        period=period, year=year, month=month, week=week,
        teacher_id=teacher_id, auth_user=auth_user
    )
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=0.5*inch, leftMargin=0.5*inch, topMargin=0.5*inch, bottomMargin=0.5*inch)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#064e3b'),
        spaceAfter=12,
        alignment=TA_CENTER
    )
    
    elements = []
    
    # Title
    elements.append(Paragraph("SIPGURU - Laporan Absensi Guru", title_style))
    elements.append(Spacer(1, 0.2*inch))
    
    # Period info
    period_text = f"Periode: {report['period'].capitalize()} ({report['start_date']} s/d {report['end_date']})"
    elements.append(Paragraph(period_text, styles['Normal']))
    elements.append(Spacer(1, 0.3*inch))
    
    # Summary table
    summary_data = [
        ["Total Absensi", str(report["summary"]["total"])],
        ["Tepat Waktu", str(report["summary"]["on_time"])],
        ["Terlambat", str(report["summary"]["late"])],
        ["Rate Kehadiran", f"{report['summary']['attendance_rate']}%"],
        ["Rata-rata Keterlambatan", f"{report['summary']['avg_lateness_minutes']} menit"],
    ]
    
    summary_table = Table(summary_data, colWidths=[2.5*inch, 2*inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#064e3b')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f0fdf4')),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#064e3b')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.4*inch))
    
    # Detail table
    elements.append(Paragraph("Detail Absensi", styles['Heading2']))
    elements.append(Spacer(1, 0.2*inch))
    
    table_data = [["Tanggal", "Waktu", "NIP", "Nama", "Ruangan", "Status", "Telat (mnt)"]]
    
    for record in report["records"]:
        created = record.get("created_at", "")
        date = created[:10] if created else ""
        time = created[11:16] if len(created) > 11 else ""
        table_data.append([
            date,
            time,
            record.get("nip", "")[:10],
            (record.get("teacher_name", "") or "")[:20],
            (record.get("room_name", "") or "")[:15],
            record.get("status", ""),
            str(record.get("lateness_minutes") or ""),
        ])
    
    detail_table = Table(table_data, colWidths=[0.9*inch, 0.7*inch, 1*inch, 1.5*inch, 1.2*inch, 0.9*inch, 0.8*inch])
    detail_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#064e3b')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(detail_table)
    
    # Footer
    elements.append(Spacer(1, 0.5*inch))
    footer = Paragraph(
        f"<i>Dicetak pada: {datetime.now().strftime('%d/%m/%Y %H:%M')} | SIPGURU System</i>",
        ParagraphStyle('Footer', parent=styles['Normal'], fontSize=8, textColor=colors.grey)
    )
    elements.append(footer)
    
    doc.build(elements)
    buffer.seek(0)
    
    filename = f"laporan_absensi_{report['start_date']}_to_{report['end_date']}.pdf"
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/teacher/{teacher_id}/summary")
async def get_teacher_summary(
    teacher_id: str,
    month: int = Query(None),
    year: int = Query(None),
    auth_user: dict = Depends(get_current_user),
):
    """Get attendance summary for a specific teacher."""
    if year is None:
        year = datetime.now().year
    if month is None:
        month = datetime.now().month
    
    conn = db._get_connection()
    cursor = conn.cursor()
    
    # Teacher info
    cursor.execute("SELECT id, name, nip, email FROM users WHERE id = ?", (teacher_id,))
    teacher = cursor.fetchone()
    if not teacher:
        conn.close()
        raise HTTPException(status_code=404, detail="Teacher not found")
    
    # Monthly stats
    month_str = f"{year}-{month:02d}"
    cursor.execute("""
        SELECT status, COUNT(*) as count, AVG(lateness_minutes) as avg_late
        FROM attendance_logs
        WHERE teacher_id = ? AND strftime('%Y-%m', created_at) = ?
        GROUP BY status
    """, (teacher_id, month_str))
    
    stats = {"total": 0, "on_time": 0, "late": 0, "absent": 0, "avg_lateness": 0}
    for row in cursor.fetchall():
        status = row["status"]
        if status in stats:
            stats[status] = row["count"]
        stats["total"] += row["count"]
        if status == "late" and row["avg_late"]:
            stats["avg_lateness"] = row["avg_late"]
    
    conn.close()
    
    return {
        "teacher": dict(teacher),
        "period": {"year": year, "month": month},
        "summary": stats,
        "attendance_rate": round(stats["on_time"] / max(stats["total"], 1) * 100, 2)
    }
